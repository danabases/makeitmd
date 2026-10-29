from __future__ import annotations

import argparse
import atexit
import base64
import html
import json
import os
import re
import subprocess
import tempfile
import urllib.error
import urllib.request
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses
from pathlib import Path

import fitz
import pymupdf4llm
from docx import Document

# Suppresses MuPDF's raw "format error: ..." console noise (e.g. malformed
# structure trees in tagged PDFs) — non-fatal parser diagnostics that don't
# affect the converted output.
fitz.TOOLS.mupdf_display_errors(False)


# User configuration. Command-line arguments can override these values.
# By default, Markdown is written beside each source file. Pass --o to write
# into a separate folder. The source folder is a required positional argument.
SCAN_SUBFOLDERS = True
SINGLE_FILE: str | None = None
# Regenerate a Markdown file when its source is newer than the existing
# Markdown (e.g. after a conversion fix). Files whose Markdown is already
# up to date are left alone. Pass --skip to skip every file that
# already has Markdown, regardless of freshness (useful for resuming).
OVERWRITE_EXISTING_MARKDOWN = True
TESSERACT_EXE = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
OCR_TEXT_THRESHOLD = 150
# Image files (jpg/png/heic) are transcribed with a local Ollama vision model.
# Install and start Ollama, then pull the model before converting images.
OLLAMA_BASE_URL = "http://localhost:11434"
# minicpm-v: markedly faster than qwen2.5vl:7b on CPU-only machines, at some
# cost to accuracy on dense/low-quality images.
OLLAMA_VISION_MODEL = "minicpm-v"
# On CPU-only machines a single large image (e.g. a full-resolution phone
# screenshot, ~4000+ prompt tokens once encoded) can take 10-15+ minutes just
# to evaluate, before any output is generated -- measured at roughly
# 0.2-0.5s/token with no GPU. 300s was cutting that off mid-request.
OLLAMA_TIMEOUT_SECONDS = 1800
# Ollama defaults every model to a 4096-token context window regardless of the
# model's actual capacity -- a single high-resolution image can encode to more
# image tokens than that, so the request must set num_ctx explicitly or
# large/detailed images fail with "exceeds the available context size".
OLLAMA_NUM_CTX = 8192
VISION_TRANSCRIPTION_PROMPT = (
    "Transcribe every piece of readable text in this image exactly as it "
    "appears, preserving layout cues such as speaker names, timestamps, and "
    "message order (this is often a chat or email screenshot). Output only "
    "the transcription in plain text or Markdown -- no commentary, no preamble."
)
DIRECT_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
HEIC_IMAGE_EXTENSIONS = {".heic", ".heif"}
IMAGE_EXTENSIONS = {*DIRECT_IMAGE_EXTENSIONS, *HEIC_IMAGE_EXTENSIONS}
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".xlsx", ".txt", ".eml", *IMAGE_EXTENSIONS}
WORD_FORMAT_XML_DOCUMENT = 16  # wdFormatXMLDocument

def markdown_filename(path: Path) -> str:
    return f"{path.name}.md"


def resolve_single_file(source_dir: Path, file_name: str, scan_subfolders: bool) -> Path:
    candidate = Path(file_name)
    if candidate.is_absolute():
        if candidate.is_file():
            return candidate
        raise FileNotFoundError(f"Could not find {file_name!r}")

    source_candidate = source_dir / candidate
    working_candidate = Path.cwd() / candidate
    source_exists = source_candidate.is_file()
    working_exists = working_candidate.is_file()
    if source_exists and working_exists:
        if source_candidate.resolve() != working_candidate.resolve():
            raise ValueError(
                f"{file_name!r} identifies files under both {source_dir} and "
                "the working directory. Use an absolute path."
            )
    if source_exists:
        return source_candidate
    if working_exists:
        return working_candidate

    if candidate.name != file_name:
        raise FileNotFoundError(f"Could not find {file_name!r} under {source_dir} or the working directory")

    pattern = "**/*" if scan_subfolders else "*"
    matches = [
        path
        for path in source_dir.glob(pattern)
        if path.is_file() and path.name.lower() == candidate.name.lower()
    ]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        choices = "\n".join(str(path) for path in matches)
        raise ValueError(f"Multiple files match {file_name!r}:\n{choices}")

    raise FileNotFoundError(f"Could not find {file_name!r} under {source_dir}")


def iter_source_files(
    source_dir: Path,
    output_dir: Path | None,
    scan_subfolders: bool,
    extensions: set[str] = SUPPORTED_EXTENSIONS,
) -> list[Path]:
    pattern = "**/*" if scan_subfolders else "*"
    files = []
    for path in source_dir.glob(pattern):
        if not path.is_file():
            continue
        if output_dir and output_dir in path.parents:
            continue
        if not is_supported_conversion_type(path, extensions):
            continue
        files.append(path)
    return sorted(files, key=lambda item: str(item).lower())


def existing_markdown_path(
    path: Path,
    source_dir: Path,
    output_dir: Path | None,
) -> Path | None:
    candidate = output_path_for_file(path, source_dir, output_dir)
    return candidate if candidate.is_file() else None


def is_supported_conversion_type(path: Path, extensions: set[str] = SUPPORTED_EXTENSIONS) -> bool:
    return path.suffix.lower() in extensions


def docx_to_markdown(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            parts.append("")
            continue

        style = paragraph.style.name.lower() if paragraph.style else ""
        if style.startswith("heading"):
            level_match = re.search(r"\d+", style)
            level = min(int(level_match.group(0)), 6) if level_match else 2
            parts.append(f"{'#' * level} {text}")
        else:
            parts.append(text)

    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
        if not rows:
            continue
        widths = [max(3, len(cell)) for cell in rows[0]]
        header = "| " + " | ".join(rows[0]) + " |"
        separator = "| " + " | ".join("-" * width for width in widths) + " |"
        body = ["| " + " | ".join(row) + " |" for row in rows[1:]]
        parts.extend(["", header, separator, *body, ""])

    return "\n".join(parts).strip() + "\n"


_word_app = None


def get_word_app():
    global _word_app
    if _word_app is None:
        import win32com.client

        app = win32com.client.Dispatch("Word.Application")
        app.Visible = False
        app.DisplayAlerts = 0
        _word_app = app
    return _word_app


@atexit.register
def _quit_word_app() -> None:
    if _word_app is not None:
        try:
            _word_app.Quit()
        except Exception:
            pass


def doc_to_markdown(path: Path) -> str:
    try:
        word = get_word_app()
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "DOC conversion requires pywin32 and Microsoft Word. "
            "Install pywin32 with: pip install pywin32"
        ) from exc
    except Exception as exc:
        raise RuntimeError(f"Could not start Microsoft Word via COM: {exc}") from exc

    try:
        document = word.Documents.Open(str(path), ReadOnly=True)
    except Exception as exc:
        raise RuntimeError(f"Microsoft Word could not open {path.name}: {exc}") from exc

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_docx = Path(temp_dir) / f"{path.stem}.docx"
        try:
            document.SaveAs(str(temp_docx), FileFormat=WORD_FORMAT_XML_DOCUMENT)
        finally:
            document.Close(False)
        return docx_to_markdown(temp_docx)


def txt_to_markdown(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return path.read_text(encoding=encoding).strip() + "\n"
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace").strip() + "\n"


def format_email_addresses(value: str | None) -> str:
    if not value:
        return ""
    addresses = []
    for display_name, address in getaddresses([value]):
        if display_name and address:
            addresses.append(f"{display_name} <{address}>")
        else:
            addresses.append(display_name or address)
    return ", ".join(addresses)


def strip_html_to_text(value: str) -> str:
    value = re.sub(r"(?is)<(script|style).*?>.*?</\1>", "", value)
    value = re.sub(r"(?i)<br\s*/?>", "\n", value)
    value = re.sub(r"(?i)</p\s*>", "\n\n", value)
    value = re.sub(r"<[^>]+>", "", value)
    value = html.unescape(value)
    return re.sub(r"\n{3,}", "\n\n", value).strip()


def email_part_to_text(part) -> str:
    payload = part.get_payload(decode=True)
    if payload is None:
        text = part.get_content()
    else:
        charset = part.get_content_charset() or "utf-8"
        text = payload.decode(charset, errors="replace")

    if part.get_content_type() == "text/html":
        return strip_html_to_text(text)
    return text.strip()


def eml_to_markdown(path: Path) -> str:
    with path.open("rb") as file:
        message = BytesParser(policy=policy.default).parse(file)

    header_names = [
        "Subject",
        "From",
        "To",
        "Cc",
        "Bcc",
        "Date",
        "Message-ID",
    ]
    parts = [f"# {message.get('Subject') or path.stem}", ""]
    parts.append("## Headers")
    parts.append("")
    for name in header_names:
        value = message.get(name)
        if not value:
            continue
        if name in {"From", "To", "Cc", "Bcc"}:
            value = format_email_addresses(value)
        parts.append(f"- **{name}:** {value}")

    plain_bodies: list[str] = []
    html_bodies: list[str] = []
    attachments: list[str] = []

    for part in message.walk():
        if part.is_multipart():
            continue

        content_disposition = part.get_content_disposition()
        filename = part.get_filename()
        content_type = part.get_content_type()
        if content_disposition == "attachment" or filename:
            attachment_name = filename or "(unnamed attachment)"
            attachments.append(f"{attachment_name} ({content_type})")
            continue

        if content_type == "text/plain":
            plain_bodies.append(email_part_to_text(part))
        elif content_type == "text/html":
            html_bodies.append(email_part_to_text(part))

    bodies = [body for body in (plain_bodies or html_bodies) if body.strip()]
    if bodies:
        parts.extend(["", "## Body", "", "\n\n---\n\n".join(bodies)])

    if attachments:
        parts.extend(["", "## Attachments", ""])
        parts.extend(f"- {attachment}" for attachment in attachments)

    return "\n".join(parts).strip() + "\n"


def markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""

    column_count = max(len(row) for row in rows)
    normalized_rows = [
        row + [""] * (column_count - len(row))
        for row in rows
    ]
    widths = [
        max(3, *(len(row[index]) for row in normalized_rows))
        for index in range(column_count)
    ]
    header = "| " + " | ".join(normalized_rows[0]) + " |"
    separator = "| " + " | ".join("-" * width for width in widths) + " |"
    body = ["| " + " | ".join(row) + " |" for row in normalized_rows[1:]]
    return "\n".join([header, separator, *body])


def cell_to_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text.replace("\r\n", "<br>").replace("\n", "<br>").replace("|", "\\|")


def xlsx_to_markdown(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "XLSX conversion requires openpyxl. Install it with: pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []

    for worksheet in workbook.worksheets:
        rows = [
            [cell_to_text(cell) for cell in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            continue

        parts.append(f"## {worksheet.title}")
        parts.append("")
        parts.append(markdown_table(rows))
        parts.append("")

    return "\n".join(parts).strip() + "\n"


def pdf_to_markdown(path: Path) -> str:
    doc = fitz.open(path)
    chunks: list[str] = []
    # use_ocr=False: pages below OCR_TEXT_THRESHOLD are OCR'd explicitly below via
    # Tesseract; without this, pymupdf4llm runs its own internal OCR pass first
    # (printing "OCR on page.number=..." for every low-text page) whose output is
    # discarded, doubling OCR work for no benefit.
    # ignore_code=True: PDFs typeset in a monospaced font (e.g. Courier statutes)
    # are otherwise treated as code blocks, which drops the spaces between words.
    markdown_pages = pymupdf4llm.to_markdown(
        str(path), page_chunks=True, use_ocr=False, ignore_code=True
    )
    for index, page in enumerate(doc):
        embedded_text = page.get_text("text").strip()
        if len(embedded_text) >= OCR_TEXT_THRESHOLD:
            # Monospaced runs still come back wrapped in inline backticks.
            markdown_text = markdown_pages[index]["text"].replace("`", "").strip()
            if len(markdown_text) < len(embedded_text) * 0.6:
                chunks.append(embedded_text)
            else:
                chunks.append(markdown_text)
            continue

        if not TESSERACT_EXE.exists():
            chunks.append(embedded_text)
            continue

        print(f"  OCR page {index + 1}/{doc.page_count}...", flush=True)
        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / f"page-{index + 1}.png"
            pixmap = page.get_pixmap(matrix=fitz.Matrix(3, 3), alpha=False)
            pixmap.save(image_path)
            result = subprocess.run(
                [
                    str(TESSERACT_EXE),
                    str(image_path),
                    "stdout",
                    "--psm",
                    "6",
                ],
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
            )
            chunks.append(result.stdout.strip())

    return "\n\n---\n\n".join(chunk for chunk in chunks if chunk).strip() + "\n"


def heic_to_png(path: Path, output_path: Path) -> None:
    try:
        from PIL import Image
        from pillow_heif import register_heif_opener
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "HEIC conversion requires Pillow and pillow-heif. "
            "Install them with: pip install Pillow pillow-heif"
        ) from exc

    register_heif_opener()
    with Image.open(path) as image:
        image.convert("RGB").save(output_path, format="PNG")


_ollama_ready = False


def ollama_tags() -> dict:
    with urllib.request.urlopen(f"{OLLAMA_BASE_URL}/api/tags", timeout=10) as response:
        return json.loads(response.read().decode("utf-8"))


OLLAMA_CONNECTION_ERRORS = (urllib.error.URLError, TimeoutError, ConnectionError)


def ensure_ollama_ready() -> None:
    global _ollama_ready
    if _ollama_ready:
        return

    try:
        tags = ollama_tags()
    except OLLAMA_CONNECTION_ERRORS as exc:
        raise RuntimeError(
            f"Ollama is unavailable at {OLLAMA_BASE_URL}. Install it from "
            "https://ollama.com/download and start it with `ollama serve`."
        ) from exc
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"Could not read Ollama's model list: {exc}") from exc

    if not isinstance(tags, dict) or not isinstance(tags.get("models"), list):
        raise RuntimeError("Ollama returned an invalid model list.")
    available_models = {
        model["name"]
        for model in tags["models"]
        if isinstance(model, dict) and isinstance(model.get("name"), str)
    }
    if not any(
        name == OLLAMA_VISION_MODEL or name.startswith(f"{OLLAMA_VISION_MODEL}:")
        for name in available_models
    ):
        raise RuntimeError(
            f"Ollama model {OLLAMA_VISION_MODEL!r} is unavailable. "
            f"Run `ollama pull {OLLAMA_VISION_MODEL}` and try again."
        )

    _ollama_ready = True


def vision_transcribe_image(path: Path) -> str:
    encoded_image = base64.b64encode(path.read_bytes()).decode("ascii")
    payload = json.dumps(
        {
            "model": OLLAMA_VISION_MODEL,
            "prompt": VISION_TRANSCRIPTION_PROMPT,
            "images": [encoded_image],
            "stream": False,
            "options": {"num_ctx": OLLAMA_NUM_CTX},
        }
    ).encode("utf-8")

    request = urllib.request.Request(
        f"{OLLAMA_BASE_URL}/api/generate",
        data=payload,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=OLLAMA_TIMEOUT_SECONDS) as response:
            body = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(
            f"Ollama at {OLLAMA_BASE_URL} rejected the vision request for model "
            f"{OLLAMA_VISION_MODEL!r}: {detail}"
        ) from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(
            f"Could not reach Ollama at {OLLAMA_BASE_URL}. Start Ollama and run "
            f"'ollama pull {OLLAMA_VISION_MODEL}' first. ({exc})"
        ) from exc
    except (TimeoutError, ConnectionError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"Ollama image request failed: {exc}") from exc

    if not isinstance(body, dict) or not isinstance(body.get("response"), str):
        raise RuntimeError("Ollama returned an invalid image response.")
    return body["response"].strip()


def image_to_markdown(path: Path) -> str:
    ensure_ollama_ready()

    if path.suffix.lower() in HEIC_IMAGE_EXTENSIONS:
        with tempfile.TemporaryDirectory() as temp_dir:
            png_path = Path(temp_dir) / f"{path.stem}.png"
            heic_to_png(path, png_path)
            return vision_transcribe_image(png_path) + "\n"

    return vision_transcribe_image(path) + "\n"


def output_path_for_file(
    path: Path,
    source_dir: Path,
    output_dir: Path | None,
) -> Path:
    if output_dir is None:
        return path.with_name(markdown_filename(path))

    try:
        relative_parent = path.parent.relative_to(source_dir)
    except ValueError:
        relative_parent = Path()
    return output_dir / relative_parent / markdown_filename(path)


def validate_unique_destinations(
    paths: list[Path], source_dir: Path, output_dir: Path | None
) -> None:
    destinations: dict[str, list[Path]] = {}
    for path in paths:
        destination = output_path_for_file(path, source_dir, output_dir)
        key = os.path.normcase(str(destination.resolve()))
        destinations.setdefault(key, []).append(path)

    collisions = [
        (destination, sources)
        for destination, sources in destinations.items()
        if len(sources) > 1
    ]
    if collisions:
        details = "\n".join(
            f"  {destination}:\n    " + "\n    ".join(map(str, sources))
            for destination, sources in collisions
        )
        raise ValueError(f"Multiple source files have the same Markdown destination:\n{details}")


def convert_file(
    path: Path,
    source_dir: Path,
    output_dir: Path | None,
) -> Path | None:
    if output_dir and output_dir in path.parents:
        return None
    if path.suffix.lower() == ".pdf":
        markdown = pdf_to_markdown(path)
    elif path.suffix.lower() == ".docx":
        markdown = docx_to_markdown(path)
    elif path.suffix.lower() == ".doc":
        markdown = doc_to_markdown(path)
    elif path.suffix.lower() == ".xlsx":
        markdown = xlsx_to_markdown(path)
    elif path.suffix.lower() == ".txt":
        markdown = txt_to_markdown(path)
    elif path.suffix.lower() == ".eml":
        markdown = eml_to_markdown(path)
    elif path.suffix.lower() in IMAGE_EXTENSIONS:
        markdown = image_to_markdown(path)
    else:
        return None

    output_path = output_path_for_file(path, source_dir, output_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(markdown, encoding="utf-8", newline="\n")
    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert PDF, Word, spreadsheet, text, email, and image files to Markdown.",
        usage="%(prog)s SOURCE [--o PATH] [--f FILE] [--all | --top] [--skip] [--t EXT [EXT ...]] [--m MODEL]",
    )
    parser.add_argument(
        "source_dir",
        nargs="?",
        type=Path,
        metavar="SOURCE",
        help="Folder to scan.",
    )
    parser.add_argument(
        "--o",
        dest="output_dir",
        type=Path,
        default=None,
        metavar="PATH",
        help="Output folder. Default: beside each source file.",
    )
    parser.add_argument(
        "--f",
        dest="file",
        default=SINGLE_FILE,
        metavar="FILE",
        help="Convert one file by absolute path, source-relative path, working-directory path, or filename.",
    )
    scope = parser.add_mutually_exclusive_group()
    scope.add_argument(
        "--all",
        dest="scan_subfolders",
        action="store_true",
        default=SCAN_SUBFOLDERS,
        help="Include subfolders (default).",
    )
    scope.add_argument(
        "--top",
        dest="scan_subfolders",
        action="store_false",
        help="Scan only the source folder.",
    )
    parser.add_argument(
        "--skip",
        dest="overwrite_existing",
        action="store_false",
        default=OVERWRITE_EXISTING_MARKDOWN,
        help="Skip files with existing Markdown, even if the source is newer.",
    )
    parser.add_argument(
        "--t",
        dest="types",
        nargs="+",
        metavar="EXT",
        default=None,
        help=(
            "Convert only these file types, e.g. --t pdf docx. "
            "Defaults to all supported types: "
            + ", ".join(sorted(ext.lstrip(".") for ext in SUPPORTED_EXTENSIONS))
        ),
    )
    parser.add_argument(
        "--m",
        dest="vision_model",
        default=OLLAMA_VISION_MODEL,
        metavar="MODEL",
        help=f"Ollama vision model. Default: {OLLAMA_VISION_MODEL}.",
    )

    # Keep commands from earlier versions working without crowding --help.
    parser.add_argument(
        "--source-dir", dest="legacy_source_dir", type=Path, help=argparse.SUPPRESS
    )
    parser.add_argument(
        "--output-dir", dest="output_dir", type=Path,
        default=argparse.SUPPRESS, help=argparse.SUPPRESS,
    )
    parser.add_argument(
        "--file", dest="file", default=argparse.SUPPRESS, help=argparse.SUPPRESS
    )
    scope.add_argument(
        "--recursive", dest="scan_subfolders", action="store_true",
        default=argparse.SUPPRESS, help=argparse.SUPPRESS,
    )
    scope.add_argument(
        "--no-recursive", dest="scan_subfolders", action="store_false",
        default=argparse.SUPPRESS, help=argparse.SUPPRESS,
    )
    parser.add_argument(
        "--skip-existing", dest="overwrite_existing", action="store_false",
        default=argparse.SUPPRESS, help=argparse.SUPPRESS,
    )
    parser.add_argument(
        "--types", dest="types", nargs="+",
        default=argparse.SUPPRESS, help=argparse.SUPPRESS,
    )
    parser.add_argument(
        "--vision-model", dest="vision_model",
        default=argparse.SUPPRESS, help=argparse.SUPPRESS,
    )
    args = parser.parse_args()

    if args.source_dir is not None and args.legacy_source_dir is not None:
        parser.error("Give SOURCE once, either as a folder or with --source-dir.")
    args.source_dir = args.source_dir or args.legacy_source_dir
    if args.source_dir is None:
        parser.error("SOURCE folder is required")

    if args.types:
        selected_extensions = {f".{ext.lower().lstrip('.')}" for ext in args.types}
        unknown = selected_extensions - SUPPORTED_EXTENSIONS
        if unknown:
            supported = ", ".join(sorted(ext.lstrip(".") for ext in SUPPORTED_EXTENSIONS))
            parser.error(
                f"Unsupported --t value(s): {', '.join(sorted(unknown))}. "
                f"Supported types: {supported}"
            )
        args.extensions = selected_extensions
    else:
        args.extensions = SUPPORTED_EXTENSIONS

    return args


def main() -> None:
    global OLLAMA_VISION_MODEL
    args = parse_args()
    OLLAMA_VISION_MODEL = args.vision_model
    source_dir = args.source_dir.resolve()
    output_dir = args.output_dir.resolve() if args.output_dir else None
    if not source_dir.is_dir():
        raise SystemExit(f"Source directory does not exist: {source_dir}")

    try:
        if args.file:
            paths = [resolve_single_file(source_dir, args.file, args.scan_subfolders)]
        else:
            paths = iter_source_files(source_dir, output_dir, args.scan_subfolders, args.extensions)
        validate_unique_destinations(paths, source_dir, output_dir)
    except (FileNotFoundError, ValueError) as exc:
        raise SystemExit(str(exc)) from None

    converted = []
    skipped_unsupported = []
    skipped_existing_markdown = []
    failed = []
    total = len(paths)
    for position, path in enumerate(paths, start=1):
        if not is_supported_conversion_type(path, args.extensions):
            skipped_unsupported.append(path)
            continue
        existing_markdown = existing_markdown_path(path, source_dir, output_dir)
        if existing_markdown is not None:
            if not args.overwrite_existing:
                skipped_existing_markdown.append(path)
                continue
            if existing_markdown.stat().st_mtime >= path.stat().st_mtime:
                skipped_existing_markdown.append(path)
                continue
        print(f"[{position}/{total}] Converting: {path}", flush=True)
        try:
            output_path = convert_file(
                path,
                source_dir=source_dir,
                output_dir=output_dir,
            )
        except (RuntimeError, OSError, ValueError) as exc:
            failed.append((path, str(exc)))
            print(f"  Failed: {exc}", flush=True)
            continue
        if output_path is not None:
            converted.append(output_path)

    print(f"Converted {len(converted)} files.")
    if skipped_unsupported:
        print(f"Skipped {len(skipped_unsupported)} unsupported files.")
        for path in skipped_unsupported:
            print(path)
    if skipped_existing_markdown:
        print(
            f"Skipped {len(skipped_existing_markdown)} files with existing Markdown outputs."
        )
        for path in skipped_existing_markdown:
            print(path)
    if failed:
        print(f"Failed to convert {len(failed)} files.")
        for path, error in failed:
            print(f"{path}: {error}")
    for path in converted:
        print(path)
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
